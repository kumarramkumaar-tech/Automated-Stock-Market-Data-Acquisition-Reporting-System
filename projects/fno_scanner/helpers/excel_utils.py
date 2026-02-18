"""Excel read/write utilities for FNO Scanner."""
import os
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill


def append_df_to_excel(filename, df, sheet_name="FNO_Data", index=False):
    """
    Append DataFrame into an existing Excel file without duplicating headers.
    Creates the file if it doesn't exist.
    """
    folder = os.path.dirname(filename)
    if folder:
        os.makedirs(folder, exist_ok=True)

    if not os.path.exists(filename):
        df.to_excel(filename, sheet_name=sheet_name, index=index)
        return

    book = load_workbook(filename)
    if sheet_name in book.sheetnames:
        ws = book[sheet_name]
        start_row = ws.max_row + 1
        with pd.ExcelWriter(filename, engine='openpyxl', mode='a',
                            if_sheet_exists='overlay') as writer:
            df.to_excel(writer, sheet_name=sheet_name, index=index,
                        header=False, startrow=start_row)
    else:
        with pd.ExcelWriter(filename, engine='openpyxl', mode='a') as writer:
            df.to_excel(writer, sheet_name=sheet_name, index=index)


def format_fno_excel(file_path, data_sheet="FNO_Data", summary_sheet="Daily_Summary",
                     analysis_sheet="Analysis"):
    """
    Auto-format FNO Scanner Excel sheets:
    - Bold headers, center alignment
    - Auto column widths
    - Color-coded change columns (green positive, red negative)
    """
    if not os.path.exists(file_path):
        return

    try:
        wb = load_workbook(file_path)

        change_columns = [
            "Price Change", "Price Change %", "OI Change", "OI Change %",
            "IV Change", "IV Change %", "Fut Price Change", "Fut Price Change %",
            "Fut OI Change", "Fut OI Change %"
        ]

        for sheet_name in [data_sheet, summary_sheet, analysis_sheet]:
            if sheet_name not in wb.sheetnames:
                continue

            ws = wb[sheet_name]
            bold = Font(bold=True)

            # Bold + center headers
            for cell in ws[1]:
                cell.font = bold
                cell.alignment = Alignment(horizontal="center")

            # Auto column width
            for col in ws.columns:
                max_length = 0
                col_letter = col[0].column_letter
                for cell in col:
                    val = str(cell.value) if cell.value is not None else ""
                    max_length = max(max_length, len(val))
                ws.column_dimensions[col_letter].width = min(max_length + 2, 35)

            # Colorize change columns
            if sheet_name == data_sheet:
                for row in ws.iter_rows(min_row=2):
                    for cell in row:
                        header = ws.cell(row=1, column=cell.column).value
                        if header in change_columns:
                            try:
                                val = float(cell.value)
                                if val > 0:
                                    cell.font = Font(color="008000")
                                elif val < 0:
                                    cell.font = Font(color="FF0000")
                            except (ValueError, TypeError):
                                pass

        wb.save(file_path)

    except Exception as e:
        print(f"Excel formatting failed: {e}")
