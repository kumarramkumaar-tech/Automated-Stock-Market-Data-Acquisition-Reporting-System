# helpers/report_utils.py – FII/DII Activity Tracker
import pandas as pd
import os
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment
from datetime import datetime
from helpers.notify import send_telegram


def compute_net_flows(df):
    """
    Compute net FII and DII flows from buy/sell data.
    Net Flow = Buy Value - Sell Value
    """
    result = df.copy()

    for col in ["FII Buy", "FII Sell", "DII Buy", "DII Sell"]:
        if col in result.columns:
            result[col] = pd.to_numeric(result[col].astype(str).str.replace(",", ""), errors="coerce").fillna(0)

    if "FII Buy" in result.columns and "FII Sell" in result.columns:
        result["FII Net"] = result["FII Buy"] - result["FII Sell"]

    if "DII Buy" in result.columns and "DII Sell" in result.columns:
        result["DII Net"] = result["DII Buy"] - result["DII Sell"]

    return result


def update_cumulative_flow(file_path):
    """
    Calculate cumulative FII/DII flows over time.
    Writes to 'Cumulative_Flow' sheet.
    """
    if not os.path.exists(file_path):
        return

    try:
        df = pd.read_excel(file_path, sheet_name="Daily_Flow")
        if df.empty:
            return

        df = compute_net_flows(df)
        df["Date"] = pd.to_datetime(df["Date"])

        # Aggregate by date
        daily = (
            df.groupby("Date")
            .agg({
                "FII Net": "sum",
                "DII Net": "sum",
                "FII Buy": "sum",
                "FII Sell": "sum",
                "DII Buy": "sum",
                "DII Sell": "sum",
            })
            .reset_index()
            .sort_values("Date")
        )

        # Cumulative sums
        daily["FII Cumulative"] = daily["FII Net"].cumsum()
        daily["DII Cumulative"] = daily["DII Net"].cumsum()
        daily["Combined Net"] = daily["FII Net"] + daily["DII Net"]

        # Round values
        for col in daily.select_dtypes(include=["float64"]).columns:
            daily[col] = daily[col].round(2)

        with pd.ExcelWriter(file_path, engine="openpyxl", mode="a", if_sheet_exists="replace") as writer:
            daily.to_excel(writer, sheet_name="Cumulative_Flow", index=False)

    except Exception as e:
        send_telegram(f"Error computing cumulative flows: {e}")


def update_summary(file_path):
    """Creates daily summary with FII/DII net flows and sentiment."""
    if not os.path.exists(file_path):
        return

    try:
        df = pd.read_excel(file_path, sheet_name="Daily_Flow")
        if df.empty or "Date" not in df.columns:
            return

        df = compute_net_flows(df)
        df["Date"] = pd.to_datetime(df["Date"]).dt.date

        summary = (
            df.groupby("Date")
            .agg({
                "FII Net": "sum",
                "DII Net": "sum",
            })
            .reset_index()
        )

        summary["FII Net"] = summary["FII Net"].round(2)
        summary["DII Net"] = summary["DII Net"].round(2)
        summary["Combined Net"] = (summary["FII Net"] + summary["DII Net"]).round(2)

        # Sentiment classification
        def classify_sentiment(row):
            if row["FII Net"] > 0 and row["DII Net"] > 0:
                return "Strong Bullish"
            elif row["FII Net"] > 0 and row["DII Net"] < 0:
                return "FII Bullish"
            elif row["FII Net"] < 0 and row["DII Net"] > 0:
                return "DII Support"
            else:
                return "Bearish"

        summary["Sentiment"] = summary.apply(classify_sentiment, axis=1)

        with pd.ExcelWriter(file_path, engine="openpyxl", mode="a", if_sheet_exists="replace") as writer:
            summary.to_excel(writer, sheet_name="Daily_Summary", index=False)

        # Telegram message for today
        today = datetime.now().date()
        today_row = summary[summary["Date"] == today]
        if not today_row.empty:
            r = today_row.iloc[0]
            msg = (
                f"*FII/DII Daily Summary ({today})*\n"
                f"FII Net: {r['FII Net']:+.2f} Cr\n"
                f"DII Net: {r['DII Net']:+.2f} Cr\n"
                f"Combined: {r['Combined Net']:+.2f} Cr\n"
                f"Sentiment: {r['Sentiment']}"
            )
            send_telegram(msg)

        # Also update cumulative flows
        update_cumulative_flow(file_path)

    except Exception as e:
        send_telegram(f"Error updating FII/DII summary: {e}")


def format_excel(file_path):
    """Auto-format all sheets in the FII/DII Excel file."""
    if not os.path.exists(file_path):
        return

    try:
        wb = load_workbook(file_path)

        for sheet_name in ["Daily_Flow", "Daily_Summary", "Cumulative_Flow"]:
            if sheet_name not in wb.sheetnames:
                continue
            ws = wb[sheet_name]
            bold = Font(bold=True)
            for cell in ws[1]:
                cell.font = bold
                cell.alignment = Alignment(horizontal="center")

            for col in ws.columns:
                max_length = 0
                col_letter = col[0].column_letter
                for cell in col:
                    val = str(cell.value) if cell.value is not None else ""
                    max_length = max(max_length, len(val))
                ws.column_dimensions[col_letter].width = min(max_length + 2, 35)

            # Color-code net flow columns (green=positive, red=negative)
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    header = ws.cell(row=1, column=cell.column).value
                    if header and ("Net" in str(header) or "Cumulative" in str(header)):
                        try:
                            val = float(cell.value)
                            if val > 0:
                                cell.font = Font(color="008000")
                            elif val < 0:
                                cell.font = Font(color="FF0000")
                        except (ValueError, TypeError):
                            pass

        wb.save(file_path)

    except Exception as e:
        send_telegram(f"Excel formatting failed: {e}")
