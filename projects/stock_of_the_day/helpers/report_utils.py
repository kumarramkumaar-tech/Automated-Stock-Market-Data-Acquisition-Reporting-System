# helpers/report_utils.py – Stock of the Day
import pandas as pd
import numpy as np
import os
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment
from datetime import datetime
from helpers.notify import send_telegram


def compute_stock_scores(df, weights):
    """
    Score each stock based on OI change, volume spike, and price momentum.
    Returns DataFrame with scores, sorted descending.
    """
    scored = df.copy()

    # Normalize Change % as OI change proxy
    scored["Change %"] = pd.to_numeric(scored["Change %"], errors="coerce").fillna(0)
    scored["Price Change %"] = pd.to_numeric(scored["Price Change %"], errors="coerce").fillna(0)

    # Composite score: weighted combination of absolute change metrics
    scored["OI_Score"] = scored["Change %"].abs()
    scored["Price_Score"] = scored["Price Change %"].abs()

    # Normalize each score to 0-100 range
    for col in ["OI_Score", "Price_Score"]:
        max_val = scored[col].max()
        if max_val > 0:
            scored[col] = (scored[col] / max_val) * 100

    # Volume proxy: count of signals per symbol (more signals = higher activity)
    symbol_counts = scored.groupby("Symbol").size().reset_index(name="Signal_Count")
    scored = scored.merge(symbol_counts, on="Symbol", how="left")
    max_count = scored["Signal_Count"].max()
    if max_count > 0:
        scored["Volume_Score"] = (scored["Signal_Count"] / max_count) * 100
    else:
        scored["Volume_Score"] = 0

    # Final composite score
    scored["Composite_Score"] = (
        weights.get("oi_change", 0.4) * scored["OI_Score"]
        + weights.get("volume_spike", 0.3) * scored["Volume_Score"]
        + weights.get("price_momentum", 0.3) * scored["Price_Score"]
    )

    return scored.sort_values("Composite_Score", ascending=False)


def pick_stock_of_the_day(file_path, weights, top_n=5):
    """
    Reads data log, computes scores, picks top N stocks,
    writes to 'Stock_Rankings' sheet, and sends Telegram alert.
    """
    if not os.path.exists(file_path):
        return

    try:
        df = pd.read_excel(file_path, sheet_name="Data_Log")
        if df.empty:
            return

        # Filter to today's data only
        df["Fetched At"] = pd.to_datetime(df["Fetched At"])
        today = datetime.now().date()
        df_today = df[df["Fetched At"].dt.date == today]

        if df_today.empty:
            send_telegram("No data collected today for Stock of the Day ranking.")
            return

        scored = compute_stock_scores(df_today, weights)

        # Aggregate by symbol: take the best score per symbol
        ranking = (
            scored.groupby("Symbol")
            .agg({
                "Composite_Score": "max",
                "Change %": "mean",
                "Price Change %": "mean",
                "Signal_Count": "first",
                "Builtup Type": "first"
            })
            .reset_index()
            .sort_values("Composite_Score", ascending=False)
            .head(top_n)
        )

        ranking["Rank"] = range(1, len(ranking) + 1)
        ranking["Date"] = today
        ranking = ranking[["Date", "Rank", "Symbol", "Composite_Score",
                           "Change %", "Price Change %", "Signal_Count", "Builtup Type"]]

        # Write to Excel
        with pd.ExcelWriter(file_path, engine="openpyxl", mode="a", if_sheet_exists="replace") as writer:
            ranking.to_excel(writer, sheet_name="Stock_Rankings", index=False)

        # Telegram message
        top_stock = ranking.iloc[0]
        msg = (
            f"*Stock of the Day ({today})*\n\n"
            f"*#{top_stock['Symbol']}*\n"
            f"Score: {top_stock['Composite_Score']:.1f}/100\n"
            f"Avg Change %: {top_stock['Change %']:.2f}\n"
            f"Avg Price Change %: {top_stock['Price Change %']:.2f}\n"
            f"Signal Count: {int(top_stock['Signal_Count'])}\n"
            f"Buildup: {top_stock['Builtup Type']}\n\n"
            f"Top {top_n} Stocks:\n"
        )
        for _, row in ranking.iterrows():
            msg += f"{int(row['Rank'])}. {row['Symbol']} (Score: {row['Composite_Score']:.1f})\n"

        send_telegram(msg)

    except Exception as e:
        send_telegram(f"Error in Stock of the Day: {e}")


def update_summary(file_path):
    """
    Creates daily summary with per-day statistics.
    """
    if not os.path.exists(file_path):
        return

    try:
        df = pd.read_excel(file_path, sheet_name="Data_Log")
        if df.empty or "Fetched At" not in df.columns:
            return

        df["Change %"] = pd.to_numeric(df["Change %"], errors="coerce")
        df["Date"] = pd.to_datetime(df["Fetched At"]).dt.date

        summary = (
            df.groupby("Date")["Change %"]
            .agg(["count", "mean", "max", "min"])
            .reset_index()
            .rename(columns={
                "count": "Total Rows",
                "mean": "Avg Change %",
                "max": "Max Change %",
                "min": "Min Change %",
            })
        )
        summary["Avg Change %"] = summary["Avg Change %"].round(2)
        summary["Max Change %"] = summary["Max Change %"].round(2)
        summary["Min Change %"] = summary["Min Change %"].round(2)

        with pd.ExcelWriter(file_path, engine="openpyxl", mode="a", if_sheet_exists="replace") as writer:
            summary.to_excel(writer, sheet_name="Daily_Summary", index=False)

        today = datetime.now().date()
        today_row = summary[summary["Date"] == today]
        if not today_row.empty:
            r = today_row.iloc[0]
            msg = (
                f"*Stock Screener Daily Summary ({today})*\n"
                f"Total Entries: {int(r['Total Rows'])}\n"
                f"Avg Change %: {r['Avg Change %']}\n"
                f"Max Change %: {r['Max Change %']}\n"
                f"Min Change %: {r['Min Change %']}"
            )
            send_telegram(msg)

    except Exception as e:
        send_telegram(f"Error updating daily summary: {e}")


def format_excel(file_path):
    """Auto-format Data_Log, Daily_Summary, and Stock_Rankings sheets."""
    if not os.path.exists(file_path):
        return

    try:
        wb = load_workbook(file_path)

        for sheet_name in ["Data_Log", "Daily_Summary", "Stock_Rankings"]:
            if sheet_name not in wb.sheetnames:
                continue
            ws = wb[sheet_name]
            bold = Font(bold=True)
            for cell in ws[1]:
                cell.font = bold
                cell.alignment = Alignment(horizontal="center")

            for col in ws.columns:
                max_length = 0
                col_letter = col[0].column_letter
                for cell in col:
                    val = str(cell.value) if cell.value is not None else ""
                    max_length = max(max_length, len(val))
                ws.column_dimensions[col_letter].width = min(max_length + 2, 35)

            # Color-code percentage columns
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    header = ws.cell(row=1, column=cell.column).value
                    if header and "%" in str(header):
                        try:
                            val = float(cell.value)
                            if val > 0:
                                cell.font = Font(color="008000")
                            elif val < 0:
                                cell.font = Font(color="FF0000")
                        except (ValueError, TypeError):
                            pass

        wb.save(file_path)

    except Exception as e:
        send_telegram(f"Excel formatting failed: {e}")
