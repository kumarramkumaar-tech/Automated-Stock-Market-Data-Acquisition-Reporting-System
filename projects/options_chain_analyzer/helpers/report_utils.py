# helpers/report_utils.py – Options Chain Analyzer
import pandas as pd
import numpy as np
import os
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment
from datetime import datetime
from helpers.notify import send_telegram


def calculate_max_pain(df):
    """
    Calculate Max Pain strike price.
    Max Pain = strike where total losses for option writers is minimized.
    """
    if df.empty:
        return None

    strikes = df["Strike"].unique()
    pain_values = {}

    for strike in strikes:
        # For each strike, calculate total intrinsic value of all options
        call_pain = 0
        put_pain = 0

        for _, row in df.iterrows():
            s = row["Strike"]
            call_oi = pd.to_numeric(row.get("Call OI", 0), errors="coerce") or 0
            put_oi = pd.to_numeric(row.get("Put OI", 0), errors="coerce") or 0

            # Call holders' gain if underlying settles at this strike
            if strike > s:
                call_pain += (strike - s) * call_oi

            # Put holders' gain if underlying settles at this strike
            if strike < s:
                put_pain += (s - strike) * put_oi

        pain_values[strike] = call_pain + put_pain

    if not pain_values:
        return None

    return min(pain_values, key=pain_values.get)


def calculate_pcr(df):
    """
    Calculate Put-Call Ratio from OI data.
    PCR > 1 = bullish (more put writing), PCR < 1 = bearish.
    """
    if df.empty:
        return None

    total_put_oi = pd.to_numeric(df.get("Put OI", pd.Series()), errors="coerce").sum()
    total_call_oi = pd.to_numeric(df.get("Call OI", pd.Series()), errors="coerce").sum()

    if total_call_oi == 0:
        return None

    return round(total_put_oi / total_call_oi, 3)


def find_support_resistance(df):
    """
    Identify support and resistance levels from OI concentration.
    - Highest Put OI strike = Support
    - Highest Call OI strike = Resistance
    """
    if df.empty:
        return None, None

    df_copy = df.copy()
    df_copy["Call OI"] = pd.to_numeric(df_copy.get("Call OI", pd.Series()), errors="coerce").fillna(0)
    df_copy["Put OI"] = pd.to_numeric(df_copy.get("Put OI", pd.Series()), errors="coerce").fillna(0)

    resistance = df_copy.loc[df_copy["Call OI"].idxmax(), "Strike"] if df_copy["Call OI"].max() > 0 else None
    support = df_copy.loc[df_copy["Put OI"].idxmax(), "Strike"] if df_copy["Put OI"].max() > 0 else None

    return support, resistance


def analyze_chain(file_path, symbol="NIFTY"):
    """
    Full options chain analysis: Max Pain, PCR, Support/Resistance.
    Writes to 'Analysis' sheet and sends Telegram summary.
    """
    if not os.path.exists(file_path):
        return

    try:
        df = pd.read_excel(file_path, sheet_name="Chain_Log")
        if df.empty:
            return

        # Filter to latest fetch and specified symbol
        df["Fetched At"] = pd.to_datetime(df["Fetched At"])
        latest_time = df["Fetched At"].max()
        df_latest = df[df["Fetched At"] == latest_time]

        if "Symbol" in df_latest.columns:
            df_latest = df_latest[df_latest["Symbol"] == symbol]

        if df_latest.empty:
            return

        # Compute analysis metrics
        max_pain = calculate_max_pain(df_latest)
        pcr = calculate_pcr(df_latest)
        support, resistance = find_support_resistance(df_latest)

        # Build analysis result
        analysis = pd.DataFrame([{
            "Date": datetime.now().strftime("%Y-%m-%d"),
            "Time": datetime.now().strftime("%H:%M:%S"),
            "Symbol": symbol,
            "Max Pain": max_pain,
            "PCR": pcr,
            "Support (Put OI)": support,
            "Resistance (Call OI)": resistance,
            "Total Call OI": pd.to_numeric(df_latest.get("Call OI", pd.Series()), errors="coerce").sum(),
            "Total Put OI": pd.to_numeric(df_latest.get("Put OI", pd.Series()), errors="coerce").sum(),
        }])

        # Append to Analysis sheet
        from helpers.excel_utils import append_df_to_excel
        append_df_to_excel(file_path, analysis, sheet_name="Analysis")

        # Telegram alert
        sentiment = "BULLISH" if pcr and pcr > 1 else "BEARISH" if pcr and pcr < 1 else "NEUTRAL"
        msg = (
            f"*Options Chain Analysis - {symbol}*\n"
            f"Time: {datetime.now().strftime('%H:%M')}\n\n"
            f"Max Pain: {max_pain}\n"
            f"PCR: {pcr} ({sentiment})\n"
            f"Support: {support}\n"
            f"Resistance: {resistance}\n"
        )
        send_telegram(msg)

    except Exception as e:
        send_telegram(f"Options chain analysis error: {e}")


def update_summary(file_path):
    """Creates daily summary of analysis metrics over time."""
    if not os.path.exists(file_path):
        return

    try:
        df = pd.read_excel(file_path, sheet_name="Analysis")
        if df.empty:
            return

        df["PCR"] = pd.to_numeric(df["PCR"], errors="coerce")
        df["Date"] = pd.to_datetime(df["Date"]).dt.date

        summary = (
            df.groupby("Date")
            .agg({
                "PCR": ["mean", "max", "min"],
                "Max Pain": "last",
                "Support (Put OI)": "last",
                "Resistance (Call OI)": "last",
            })
            .reset_index()
        )
        summary.columns = [
            "Date", "Avg PCR", "Max PCR", "Min PCR",
            "Last Max Pain", "Last Support", "Last Resistance"
        ]

        with pd.ExcelWriter(file_path, engine="openpyxl", mode="a", if_sheet_exists="replace") as writer:
            summary.to_excel(writer, sheet_name="Daily_Summary", index=False)

        today = datetime.now().date()
        today_row = summary[summary["Date"] == today]
        if not today_row.empty:
            r = today_row.iloc[0]
            msg = (
                f"*Options Daily Summary ({today})*\n"
                f"Avg PCR: {r['Avg PCR']:.3f}\n"
                f"Last Max Pain: {r['Last Max Pain']}\n"
                f"Support: {r['Last Support']}\n"
                f"Resistance: {r['Last Resistance']}"
            )
            send_telegram(msg)

    except Exception as e:
        send_telegram(f"Error updating options summary: {e}")


def format_excel(file_path):
    """Auto-format all sheets in the options chain Excel file."""
    if not os.path.exists(file_path):
        return

    try:
        wb = load_workbook(file_path)

        for sheet_name in ["Chain_Log", "Analysis", "Daily_Summary"]:
            if sheet_name not in wb.sheetnames:
                continue
            ws = wb[sheet_name]
            bold = Font(bold=True)
            for cell in ws[1]:
                cell.font = bold
                cell.alignment = Alignment(horizontal="center")

            for col in ws.columns:
                max_length = 0
                col_letter = col[0].column_letter
                for cell in col:
                    val = str(cell.value) if cell.value is not None else ""
                    max_length = max(max_length, len(val))
                ws.column_dimensions[col_letter].width = min(max_length + 2, 35)

        wb.save(file_path)

    except Exception as e:
        send_telegram(f"Excel formatting failed: {e}")
